from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .calendar_rules import calcular_vencimento_mensal, competencia_para_texto
from .models import ItemPlanilha, ResultadoPlanilha


CABECALHOS_OBRIGATORIOS = {
    "Codigo",
    "Tipo Inscrição",
    "Servico",
    "TAG",
    "Inscrição",
    "FGTS",
    "FGTS Aprendiz",
}


def _digitos(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\D", "", str(valor))


def normalizar_inscricao(tipo: str, valor: Any) -> tuple[str, str | None]:
    tipo_norm = str(tipo or "").strip().upper()
    digitos = _digitos(valor)

    if tipo_norm == "CNPJ":
        if not digitos:
            return "", "CNPJ vazio"
        digitos = digitos.zfill(14)
        if len(digitos) != 14:
            return digitos, f"CNPJ deve ter 14 dígitos, encontrado {len(digitos)}"
        return digitos, None

    if tipo_norm == "CNO":
        if not digitos:
            return "", "CNO vazio"
        if len(digitos) != 12:
            return digitos, f"CNO deve ter 12 dígitos, encontrado {len(digitos)}"
        return digitos, None

    return digitos, f"Tipo de inscrição inválido: {tipo_norm or '(vazio)'}"


def _como_float(valor: Any) -> float | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _como_data(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


def ler_planilha(caminho: str | Path) -> ResultadoPlanilha:
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(caminho)

    wb = load_workbook(caminho, data_only=True, read_only=False)
    ws = wb["fgts_servicos (1)"] if "fgts_servicos (1)" in wb.sheetnames else wb.active

    competencia_data = _como_data(ws["B1"].value)
    vencimento_conferencia = _como_data(ws["E1"].value)

    if competencia_data is None:
        competencia = ""
        ano = mes = 0
    else:
        competencia = competencia_para_texto(competencia_data)
        ano, mes = competencia_data.year, competencia_data.month

    vencimento_calculado = calcular_vencimento_mensal(ano, mes) if ano and mes else date.min

    resultado = ResultadoPlanilha(
        caminho=str(caminho),
        aba=ws.title,
        competencia=competencia,
        vencimento_calculado=vencimento_calculado,
        vencimento_conferencia=vencimento_conferencia,
    )

    if competencia_data is None:
        resultado.erros.append("B1 não contém uma competência reconhecível como data do Excel.")

    cabecalhos = {str(ws.cell(3, col).value or "").strip(): col for col in range(1, ws.max_column + 1)}
    faltantes = CABECALHOS_OBRIGATORIOS - set(cabecalhos)
    if faltantes:
        resultado.erros.append("Cabeçalhos ausentes: " + ", ".join(sorted(faltantes)))
        return resultado

    vistos_codigo: set[str] = set()
    vistos_inscricao: set[str] = set()
    vistos_tag: set[str] = set()

    for linha in range(4, ws.max_row + 1):
        valores = {nome: ws.cell(linha, coluna).value for nome, coluna in cabecalhos.items()}
        if all(valores.get(nome) in (None, "") for nome in CABECALHOS_OBRIGATORIOS):
            continue

        codigo = str(valores.get("Codigo") or "").strip()
        tipo = str(valores.get("Tipo Inscrição") or "").strip().upper()
        servico = str(valores.get("Servico") or "").strip()
        tag = str(valores.get("TAG") or "").strip()
        inscricao, erro_inscricao = normalizar_inscricao(tipo, valores.get("Inscrição"))

        erros_linha: list[str] = []
        if not codigo:
            erros_linha.append("Codigo vazio")
        if not servico:
            erros_linha.append("Servico vazio")
        if not tag:
            erros_linha.append("TAG vazia")
        if erro_inscricao:
            erros_linha.append(erro_inscricao)

        if codigo and codigo in vistos_codigo:
            erros_linha.append(f"Codigo duplicado: {codigo}")
        if inscricao and inscricao in vistos_inscricao:
            erros_linha.append(f"Inscrição duplicada: {inscricao}")
        if tag and tag in vistos_tag:
            erros_linha.append(f"TAG duplicada: {tag}")

        if erros_linha:
            resultado.erros.append(f"Linha {linha}: " + "; ".join(erros_linha))
            continue

        vistos_codigo.add(codigo)
        vistos_inscricao.add(inscricao)
        vistos_tag.add(tag)

        resultado.itens.append(
            ItemPlanilha(
                linha=linha,
                codigo=codigo,
                tipo_inscricao=tipo,
                servico=servico,
                tag=tag,
                inscricao=inscricao,
                fgts_referencia=_como_float(valores.get("FGTS")),
                fgts_aprendiz_referencia=_como_float(valores.get("FGTS Aprendiz")),
            )
        )

    if resultado.vencimento_conferencia is None:
        resultado.avisos.append("E1 não contém vencimento de conferência reconhecível como data.")
    elif resultado.vencimento_conferencia != resultado.vencimento_calculado:
        resultado.avisos.append(
            "Vencimento divergente: calculado "
            f"{resultado.vencimento_calculado:%d/%m/%Y}, planilha {resultado.vencimento_conferencia:%d/%m/%Y}."
        )

    if not resultado.itens and not resultado.erros:
        resultado.erros.append("Nenhum registro válido encontrado a partir da linha 4.")

    return resultado
