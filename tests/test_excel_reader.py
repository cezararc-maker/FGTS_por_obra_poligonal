import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from fgts_obra.excel_reader import ler_planilha, normalizar_inscricao, normalizar_tag


class TestExcelReader(unittest.TestCase):
    def test_tag_remove_caracteres_recusados_e_preserva_permitidos(self):
        self.assertEqual(
            normalizar_tag("86-ALMOXARIFADO HU UFGD(EBS) #1"),
            "86-ALMOXARIFADO HU UFGDEBS 1",
        )
        self.assertEqual(normalizar_tag("obra áçú $:%*+-./"), "OBRA ACU $:%*+-./")

    def test_importacao_usa_tag_limpa_e_avisa(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "tag.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["B1"] = date(2026, 8, 1)
            ws["E1"] = date(2026, 9, 18)
            self._preencher_registro_minimo(ws)
            ws.cell(4, 4).value = "86-ALMOXARIFADO HU UFGD(EBS)"
            wb.save(caminho)
            resultado = ler_planilha(caminho)
            self.assertTrue(resultado.valida)
            self.assertEqual(resultado.itens[0].tag, "86-ALMOXARIFADO HU UFGDEBS")
            self.assertTrue(any("TAG ajustada" in aviso for aviso in resultado.avisos))
            self.assertEqual(ws.cell(4, 4).value, "86-ALMOXARIFADO HU UFGD(EBS)")

    def test_tag_vazia_ou_longa_apos_limpeza_invalida_planilha(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "tag_invalida.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["B1"] = date(2026, 8, 1)
            ws["E1"] = date(2026, 9, 18)
            self._preencher_registro_minimo(ws)
            ws.cell(4, 4).value = "()#"
            wb.save(caminho)
            self.assertIn("TAG vazia após remover", " ".join(ler_planilha(caminho).erros))
            ws.cell(4, 4).value = "A" * 41
            wb.save(caminho)
            self.assertIn("excede 40", " ".join(ler_planilha(caminho).erros))

    def test_cnpj_com_zero_inicial_e_reconstruido(self):
        inscricao, erro = normalizar_inscricao("CNPJ", 3492162000182)
        self.assertIsNone(erro)
        self.assertEqual(len(inscricao), 14)
        self.assertTrue(inscricao.startswith("0"))

    def test_cno_deve_ter_12_digitos(self):
        inscricao, erro = normalizar_inscricao("CNO", "900000000000")
        self.assertIsNone(erro)
        self.assertEqual(inscricao, "900000000000")

    def _preencher_registro_minimo(self, ws) -> None:
        cabecalhos = [
            "Codigo",
            "Tipo Inscrição",
            "Servico",
            "TAG",
            "Inscrição",
            "FGTS",
            "FGTS Aprendiz",
        ]
        for coluna, nome in enumerate(cabecalhos, start=1):
            ws.cell(3, coluna).value = nome

        ws.cell(4, 1).value = "001"
        ws.cell(4, 2).value = "CNO"
        ws.cell(4, 3).value = "OBRA TESTE"
        ws.cell(4, 4).value = "1 - OBRA TESTE"
        ws.cell(4, 5).value = "900000000000"
        ws.cell(4, 6).value = 100.0
        ws.cell(4, 7).value = 0.0

    def test_le_planilha_minima_valida(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "modelo.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "fgts_servicos (1)"
            ws["B1"] = date(2026, 8, 1)
            ws["E1"] = date(2026, 9, 18)
            self._preencher_registro_minimo(ws)
            wb.save(caminho)

            resultado = ler_planilha(caminho)
            self.assertTrue(resultado.valida)
            self.assertEqual(resultado.competencia, "08/2026")
            self.assertEqual(resultado.vencimento_calculado, date(2026, 9, 18))
            self.assertEqual(len(resultado.itens), 1)

    def test_aceita_competencia_e_vencimento_como_texto(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "modelo_texto.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "FGTS por Obra"
            ws["B1"] = "08/2026"
            ws["E1"] = "18/09/2026"
            self._preencher_registro_minimo(ws)
            wb.save(caminho)

            resultado = ler_planilha(caminho)
            self.assertTrue(resultado.valida)
            self.assertEqual(resultado.competencia, "08/2026")
            self.assertEqual(resultado.vencimento_conferencia, date(2026, 9, 18))
            self.assertEqual(resultado.vencimento_calculado, date(2026, 9, 18))
            self.assertEqual(len(resultado.itens), 1)


if __name__ == "__main__":
    unittest.main()
