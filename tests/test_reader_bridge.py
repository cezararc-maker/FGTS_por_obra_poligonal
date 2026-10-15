import io
import json
import tempfile
import threading
import unittest
from datetime import date
from http.server import HTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from unittest.mock import Mock, patch

from openpyxl import Workbook

from fgts_obra import reader_bridge
from fgts_obra.app_final import main


def workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws["B1"] = date(2026, 9, 1)
    ws["E1"] = date(2026, 10, 20)
    for col, value in enumerate(("Codigo", "Tipo Inscrição", "Servico", "TAG", "Inscrição", "FGTS", "FGTS Aprendiz"), 1):
        ws.cell(3, col).value = value
    for col, value in enumerate(("26", "CNO", "OBRA", "26-OBRA 900153138674", "900153138674", 100, 0), 1):
        ws.cell(4, col).value = value
    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


class ReaderBridgeTests(unittest.TestCase):
    def test_saves_and_validates_generated_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            path = reader_bridge._save_imported_xlsx(workbook_bytes(), Path(folder))
            self.assertTrue(path.is_file())
            self.assertIn("092026", path.name)
            self.assertEqual(path.read_bytes()[:2], b"PK")

    def test_http_import_requires_token_and_opens_guide(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            html = folder / "Leitor_FGTS.html"
            html.write_text("<html>leitor</html>", encoding="utf-8")
            launched = []
            server = HTTPServer(("127.0.0.1", 0), reader_bridge._handler(html, "secret", launched.append, folder))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                origin = f"http://127.0.0.1:{server.server_port}"
                with urlopen(f"{origin}/?token=secret") as response:
                    self.assertIn(b"leitor", response.read())
                request = Request(f"{origin}/api/import", data=workbook_bytes(), method="POST", headers={
                    "Origin": origin,
                    "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "X-FGTS-Token": "secret",
                })
                with urlopen(request) as response:
                    result = json.load(response)
                self.assertTrue(result["ok"])
                self.assertEqual(len(launched), 1)
                self.assertTrue(launched[0].is_file())
                bad = Request(f"{origin}/api/import", data=workbook_bytes(), method="POST", headers={
                    "Origin": origin,
                    "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "X-FGTS-Token": "wrong",
                })
                with self.assertRaises(HTTPError) as error:
                    urlopen(bad)
                self.assertEqual(error.exception.code, 403)
                self.assertEqual(len(launched), 1)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=2)

    @patch("fgts_obra.app_final.AppFGTSFinal")
    def test_guide_preloads_imported_path_without_starting_lot(self, app_class):
        app = app_class.return_value
        app.after.side_effect = lambda _delay, callback: callback()
        main(["--planilha", "C:/temp/FGTS_Extrator_092026.xlsx"])
        app._validar_planilha.assert_called_once()
        app._iniciar_lote.assert_not_called()
        app.mainloop.assert_called_once()


if __name__ == "__main__":
    unittest.main()
